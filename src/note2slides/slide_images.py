"""プレゼンテーション資料をスライド 1 枚ずつの画像に変換する。

    .pptx / .odp --(LibreOffice)--> PDF --(pypdfium2)--> slide_001.png ...

PDF を経由するのは、スライドの見た目をベクタのまま受け取り、目的の解像度で
ラスタライズするため。拡大による劣化が起きない。

出力は動画生成でそのまま使えるように、全スライドで同じ画素数・アルファなしに
そろえる。元のスライドが 16:9 でない場合だけ、背景色で余白を足して合わせる。
"""

from __future__ import annotations

import json
import os
import re
import shutil
import tempfile
from dataclasses import dataclass, field
from typing import Callable, List, Optional, Tuple

from . import soffice as soffice_mod

MANIFEST_NAME = "slides.json"

#: 動画側で扱いやすい形式のみを対象にする。
FORMATS = {
    "png": ("PNG", ".png"),
    "jpg": ("JPEG", ".jpg"),
}

#: LibreOffice で PDF にしてから読み込む拡張子。
PRESENTATION_SUFFIXES = (".pptx", ".ppt", ".ppsx", ".pps", ".odp")
PDF_SUFFIXES = (".pdf",)


class SlideImageError(RuntimeError):
    """スライド画像の書き出しに失敗した場合。"""


class OutputExistsError(SlideImageError):
    """出力先に前回の画像が残っている場合(--force で上書きできる)。"""


@dataclass
class ImageOptions:
    """出力画像の指定。既定は 1920x1080 (16:9) の PNG。"""

    width: int = 1920
    height: Optional[int] = None  # 未指定なら width から 16:9 で決める
    fmt: str = "png"
    jpeg_quality: int = 92
    prefix: str = "slide_"
    digits: int = 3
    background: Tuple[int, int, int] = (255, 255, 255)

    def resolved_height(self) -> int:
        if self.height:
            return self.height
        return round(self.width * 9 / 16)

    def size(self) -> Tuple[int, int]:
        return self.width, self.resolved_height()

    def suffix(self) -> str:
        return FORMATS[self.fmt][1]

    def filename(self, index: int) -> str:
        return f"{self.prefix}{index:0{self.digits}d}{self.suffix()}"

    def validate(self) -> None:
        if self.fmt not in FORMATS:
            known = " / ".join(sorted(FORMATS))
            raise SlideImageError(f"未対応の画像形式です: {self.fmt}(利用できるのは {known})")
        width, height = self.size()
        if width < 16 or height < 16:
            raise SlideImageError(f"画像サイズが小さすぎます: {width}x{height}")
        # H.264 などは偶数の画素数を前提にするため、奇数だと動画化でつまずく。
        if width % 2 or height % 2:
            raise SlideImageError(
                f"画像サイズは偶数にしてください(動画化で扱えません): {width}x{height}"
            )
        if self.digits < 1:
            raise SlideImageError("連番の桁数は 1 以上にしてください")


@dataclass
class SlideImage:
    index: int  # 1 起点。スライドの並び順。
    path: str
    width: int
    height: int

    @property
    def filename(self) -> str:
        return os.path.basename(self.path)


@dataclass
class ExportResult:
    images: List[SlideImage] = field(default_factory=list)
    source: str = ""
    pdf_path: Optional[str] = None  # --keep-pdf を指定した場合のみ残る
    manifest_path: Optional[str] = None
    warnings: List[str] = field(default_factory=list)

    @property
    def count(self) -> int:
        return len(self.images)


def export_slide_images(
    source: str,
    outdir: str,
    options: Optional[ImageOptions] = None,
    soffice_path: Optional[str] = None,
    force: bool = False,
    keep_pdf: bool = False,
    timeout: float = 180,
    on_progress: Optional[Callable[[SlideImage], None]] = None,
) -> ExportResult:
    """資料の各スライドを画像として outdir に書き出す。

    ファイル名は `slide_001.png` のように 1 起点の連番で、辞書順とスライド順が
    一致する。あわせて構成を `slides.json` に書き出す。
    """
    options = options or ImageOptions()
    options.validate()

    if not os.path.isfile(source):
        raise SlideImageError(f"入力ファイルが見つかりません: {source}")
    suffix = os.path.splitext(source)[1].lower()
    if suffix not in PRESENTATION_SUFFIXES + PDF_SUFFIXES:
        known = " / ".join(PRESENTATION_SUFFIXES + PDF_SUFFIXES)
        raise SlideImageError(f"未対応の入力形式です: {suffix or source}(扱えるのは {known})")

    os.makedirs(outdir, exist_ok=True)
    existing = _existing_outputs(outdir, options)
    if existing and not force:
        raise OutputExistsError(
            f"出力先に既に画像があります: {outdir}({len(existing)} 件)\n"
            "上書きする場合は --force を付けてください。"
        )

    result = ExportResult(source=os.path.abspath(source))
    expected = _expected_slide_count(source)

    with tempfile.TemporaryDirectory(prefix="note2slides-") as workdir:
        if suffix in PDF_SUFFIXES:
            pdf_path = os.path.abspath(source)
        else:
            pdf_path = soffice_mod.convert_to_pdf(
                source, workdir, soffice=soffice_path, timeout=timeout
            )

        result.images = _render_pdf(pdf_path, outdir, options, on_progress)
        if keep_pdf and pdf_path != os.path.abspath(source):
            result.pdf_path = _copy_pdf(pdf_path, outdir, source)

    if not result.images:
        raise SlideImageError(f"スライドが 1 枚もありませんでした: {source}")
    if expected is not None and expected != result.count:
        result.warnings.append(
            f"資料のスライド数({expected})と出力した画像の枚数({result.count})が一致しません。"
            "非表示スライドや変換時の脱落がないか確認してください。"
        )

    # 前回の残骸が次の処理に混ざらないよう、今回書かなかった分を消す。
    stale = sorted(set(existing) - {img.path for img in result.images})
    for path in stale:
        os.remove(path)
    if stale:
        result.warnings.append(f"前回の画像 {len(stale)} 件を削除しました。")

    result.manifest_path = _write_manifest(outdir, options, result)
    return result


def render_pdf_to_images(
    pdf_path: str,
    outdir: str,
    options: Optional[ImageOptions] = None,
    on_progress: Optional[Callable[[SlideImage], None]] = None,
) -> List[SlideImage]:
    """PDF の各ページを画像として書き出す(LibreOffice を使わない経路)。"""
    options = options or ImageOptions()
    options.validate()
    os.makedirs(outdir, exist_ok=True)
    return _render_pdf(pdf_path, outdir, options, on_progress)


# ---------------------------------------------------------------------------
# 内部
# ---------------------------------------------------------------------------


def _render_pdf(
    pdf_path: str,
    outdir: str,
    options: ImageOptions,
    on_progress: Optional[Callable[[SlideImage], None]],
) -> List[SlideImage]:
    pdfium = _import_pdfium()
    pil_format = FORMATS[options.fmt][0]
    target = options.size()

    try:
        document = pdfium.PdfDocument(pdf_path)
    except Exception as exc:  # pypdfium2 は独自の例外を投げる
        raise SlideImageError(f"PDF を開けませんでした: {pdf_path}\n{exc}") from exc

    images: List[SlideImage] = []
    try:
        for number, page in enumerate(document, start=1):
            try:
                image = _render_page(page, target, options.background)
            except Exception as exc:
                raise SlideImageError(
                    f"{number} 枚目のスライドを画像にできませんでした: {pdf_path}\n{exc}"
                ) from exc
            path = os.path.join(outdir, options.filename(number))
            save_options = {"quality": options.jpeg_quality} if pil_format == "JPEG" else {}
            image.save(path, format=pil_format, **save_options)
            slide_image = SlideImage(number, os.path.abspath(path), target[0], target[1])
            images.append(slide_image)
            if on_progress:
                on_progress(slide_image)
    finally:
        document.close()
    return images


def _render_page(page, target: Tuple[int, int], background: Tuple[int, int, int]):
    """1 ページを目的の画素数に収めて描画する。

    縦横比が違う場合だけ背景色で余白を足す。全スライドの画素数をそろえないと、
    後続の動画生成でフレームサイズが揃わない。
    """
    from PIL import Image

    target_width, target_height = target
    page_width = page.get_width()
    page_height = page.get_height()
    scale = min(target_width / page_width, target_height / page_height)
    rendered = page.render(scale=scale).to_pil().convert("RGB")

    if rendered.size == target:
        return rendered
    if rendered.width > target_width or rendered.height > target_height:
        # 端数で 1px はみ出すことがある。縮めて収める。
        rendered.thumbnail(target, Image.LANCZOS)
    canvas = Image.new("RGB", target, background)
    canvas.paste(rendered, ((target_width - rendered.width) // 2, (target_height - rendered.height) // 2))
    return canvas


def _import_pdfium():
    try:
        import pypdfium2
    except ImportError as exc:  # pragma: no cover - 依存が入っていない環境向け
        raise SlideImageError(
            "pypdfium2 が見つかりません。`pip install -e .` で依存を入れてください。"
        ) from exc
    return pypdfium2


def _expected_slide_count(source: str) -> Optional[int]:
    """.pptx なら表示対象のスライド数を返す。判断できない形式では None。

    LibreOffice は壊れた .pptx をテキストとして読み込み、それらしい PDF を
    出してしまうことがある。ここで先に開いて弾く。
    """
    if os.path.splitext(source)[1].lower() != ".pptx":
        return None

    from pptx import Presentation
    from pptx.exc import PythonPptxError

    try:
        presentation = Presentation(source)
        # 非表示スライドは PDF に出力されないため、あらかじめ除いて数える。
        return sum(
            1 for slide in presentation.slides if slide.element.get("show") != "0"
        )
    except (PythonPptxError, ValueError, KeyError, OSError) as exc:
        raise SlideImageError(
            f"資料として読み取れませんでした: {source}\n{type(exc).__name__}: {exc}"
        ) from exc


def _existing_outputs(outdir: str, options: ImageOptions) -> List[str]:
    """同じ命名規則で書かれた既存ファイルを列挙する。

    形式を変えて出し直した場合(png -> jpg)も前回分を見落とさないよう、
    拡張子は扱えるものすべてを対象にする。
    """
    suffixes = "|".join(re.escape(suffix) for _, suffix in FORMATS.values())
    pattern = re.compile(re.escape(options.prefix) + r"\d+(" + suffixes + r")$")
    return [
        os.path.abspath(os.path.join(outdir, name))
        for name in sorted(os.listdir(outdir))
        if pattern.match(name)
    ]


def _copy_pdf(pdf_path: str, outdir: str, source: str) -> str:
    destination = os.path.join(
        outdir, os.path.splitext(os.path.basename(source))[0] + ".pdf"
    )
    shutil.copyfile(pdf_path, destination)
    return os.path.abspath(destination)


def _write_manifest(outdir: str, options: ImageOptions, result: ExportResult) -> str:
    """後続の動画生成が読む一覧。パスは outdir からの相対で書く。"""
    width, height = options.size()
    data = {
        "source": os.path.basename(result.source),
        "format": options.fmt,
        "width": width,
        "height": height,
        "count": result.count,
        "slides": [
            {"index": image.index, "file": image.filename} for image in result.images
        ],
        "warnings": result.warnings,
    }
    path = os.path.join(outdir, MANIFEST_NAME)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write("\n")
    return os.path.abspath(path)


def ffmpeg_pattern(options: ImageOptions) -> str:
    """ffmpeg の image2 で読み込むときのパターン。"""
    return f"{options.prefix}%0{options.digits}d{options.suffix()}"
