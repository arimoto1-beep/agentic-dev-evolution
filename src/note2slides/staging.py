"""1 枚の画面を、ナレーションに合わせて段階的に見せるための「段」を作る。

    資料(.pptx) + スライド画像 + 原稿の合図 --> 段ごとの画像(steps/)

資料を画像にして音声を付けただけの動画は、1 枚の画面が **出揃った状態で
数十秒止まる**。話は進んでいるのに画面が変わらないので、見ている側は
「読み上げが終わるのを待つ」ことになる(gen36。22 分の動画で画面が変わるのが
40 回、止まったままの長さの中央値が 32 秒だった)。

ここでは 1 枚を「段」に分ける。段は **画面に置かれた順** に数える。

    表            見出しの行はそのまま、本文の 1 行が 1 段
    図解(流れ・枠・階段)  項目 1 つが 1 段(前の矢印や蹴込みは、その項目に付く)
    文章          段落 1 つが 1 段
    絵            1 段。ただし `図@1.png` `図@2.png` … が隣にあれば、その枚数ぶんの段
    それ以外の図解  図解 1 つが 1 段

n 段目の画面は、「n 段目までを出し、n 段目だけを濃く、それより前は少し引く」。
絵の段は、絵を描いた側が用意した n 枚目をそのまま出す(何を足し、何を
目立たせるかは絵ごとに違うので、ここでは決めない)。`図@2/000.png …` が
あれば、2 段目に入るときにその連番を動きとして流す。

原稿の合図(`〔2〕` `〔全体〕`。`reading.py`)が、どの文から何段目にするかを決める。
合図が無い画面には何もしない(これまでどおり 1 枚の絵のまま映る)。段の合図がある
画面は、**最初の合図までは何も出さずに**(見出しと札だけで)始まる。出揃った画面を
先に見せてから隠すと、これから話すことを先に読まれてしまう。最初から出しておきたい
ときは、最初の文に合図を書く。

資料は作り直さない。**出来上がったスライド画像の上で** 段を作る。図形の場所は
資料から読み、表の行と文章の段落だけは、場所が資料に書かれていない(描く側が
決める)ので画像から読み取る。読み取れなかったときは、その表・文章をまとめて
1 段として扱い、警告する(黙って違う場所を隠さない)。
"""

from __future__ import annotations

import json
import os
import re
import shutil
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
from PIL import Image

from . import slide_images as images_mod
from .model import (
    DIAGRAM_FLOW,
    DIAGRAM_FLOW_ACROSS,
    DIAGRAM_FRAME,
    DIAGRAM_STEPS,
    SHAPE_DIAGRAM,
    SHAPE_DIAGRAM_ITEM,
    SHAPE_FOOTER,
    SHAPE_LABEL,
    SHAPE_TABLE,
    parse_shape_name,
)
from .reading import CUE_ALL

#: 何も出していない状態(見出しと札だけ)を指す合図。
CUE_START = "0"

STEPS_DIR = "steps"
MANIFEST_NAME = "steps.json"

#: まだ出していない段を、地の色へ寄せる割合(1.0 = 見えない)。
HIDDEN = 1.0
#: 話し終えた段を、地の色へ寄せる割合。読める濃さは残す(見比べるため)。
PAST = 0.45
#: 隠すときに、図形の縁(線の太さとにじみ)まで覆うために広げる幅(画素)。
#: 広げすぎると、隣の図形の縁を削る(階段図の蹴込みは、上の段の底に接している)。
_PAD = 2

Box = Tuple[int, int, int, int]  # 左・上・右・下(画素)


class StagingError(RuntimeError):
    """段を作れなかった場合(合図が指す段が無い、など)。"""


@dataclass
class Element:
    """画面の上で、まとめて出したり引いたりする 1 つのもの。"""

    box: Box
    kind: str = "item"  # item / figure
    #: まとまりを作っている図形ごとの場所。隠す・引くのは、この 1 つずつに対して行う
    #: (まとめた四角で覆うと、あいだにある別の図形まで消える)。空なら `box` 全体。
    parts: List[Box] = field(default_factory=list)
    #: この段と一緒に現れる入れ物(枠図の外枠)。この段を隠すときだけ一緒に隠し、
    #: 薄くするときは触らない(中には、ほかの段も入っている)。
    frame: Optional[Box] = None
    frames: List[str] = field(default_factory=list)  # 絵の段(`図@n.png`)
    final: str = ""  # 絵の最後の姿(`図.png`)
    motions: Dict[int, List[str]] = field(default_factory=dict)  # 段 -> 動きの連番

    @property
    def count(self) -> int:
        return max(1, len(self.frames))


@dataclass
class StagedSlide:
    index: int
    count: int = 0  # 段の数
    images: Dict[str, str] = field(default_factory=dict)  # 合図 -> 画像
    motions: Dict[str, List[str]] = field(default_factory=dict)  # 合図 -> 動きの連番

    def to_dict(self) -> dict:
        return {
            "index": self.index,
            "count": self.count,
            "images": {key: os.path.basename(path) for key, path in self.images.items()},
            "motions": {
                key: [os.path.relpath(path, os.path.dirname(os.path.dirname(path))).replace(os.sep, "/")
                      for path in paths]
                for key, paths in self.motions.items()
            },
        }


@dataclass
class StagingResult:
    outdir: str = ""
    slides: List[StagedSlide] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    manifest_path: Optional[str] = None


# ---------------------------------------------------------------------------
# 段を作る
# ---------------------------------------------------------------------------


def build_steps(
    source: str,
    slides_dir: str,
    wanted: Dict[int, Sequence[str]],
    figures_dir: Optional[str] = None,
) -> StagingResult:
    """合図のある画面について、段ごとの画像を `slides_dir/steps/` に書き出す。

    `wanted` はスライド番号 -> そのスライドの原稿に書かれた合図(`2` や `全体`)。
    合図が指す段が画面に無いときは、動画にせずに止める(違う段が出た動画は、
    通しで見るまで気付けない)。
    """
    from pptx import Presentation

    outdir = os.path.join(slides_dir, STEPS_DIR)
    if os.path.isdir(outdir):
        shutil.rmtree(outdir)  # 前回の段が、作り直した資料に混ざらないようにする
    result = StagingResult(outdir=os.path.abspath(outdir))
    wanted = {index: list(cues) for index, cues in wanted.items() if cues}
    if not wanted:
        return result
    os.makedirs(outdir, exist_ok=True)

    images = _slide_image_paths(slides_dir)
    presentation = Presentation(source)
    width_emu = presentation.slide_width
    base_dir = os.path.dirname(os.path.abspath(source))
    number = 0
    for slide in presentation.slides:
        if slide.element.get("show") == "0":
            continue
        number += 1
        if number not in wanted:
            continue
        if number not in images:
            raise StagingError(f"スライド {number} の画像がありません: {slides_dir}")
        with Image.open(images[number]) as opened:
            base = opened.convert("RGB")
        scale = base.width / width_emu
        elements, warnings = slide_elements(
            slide, base, scale, [base_dir] + ([figures_dir] if figures_dir else [])
        )
        result.warnings.extend(f"スライド {number}: {w}" for w in warnings)
        result.slides.append(
            _stage_slide(number, base, elements, wanted[number], outdir, _title_of(slide))
        )

    result.manifest_path = _write_manifest(result)
    return result


def _stage_slide(
    number: int,
    base: Image.Image,
    elements: List[Element],
    cues: Sequence[str],
    outdir: str,
    title: str,
) -> StagedSlide:
    total = sum(element.count for element in elements)
    staged = StagedSlide(index=number, count=total)
    has_figure_steps = any(element.frames for element in elements)
    wanted = list(dict.fromkeys(cues))  # 書かれた順のまま、重複を除く
    if any(cue != CUE_ALL for cue in wanted) and CUE_START not in wanted:
        wanted.insert(0, CUE_START)  # 最初の合図までの画面(何も出していない状態)
    for cue in wanted:
        if cue == CUE_ALL:
            if not has_figure_steps:
                continue  # スライド画像そのものが「全体」
            image = render_step(base, elements, None)
            path = os.path.join(outdir, f"slide_{number:03d}.all.png")
        else:
            step = int(cue)
            if not 0 <= step <= total:
                raise StagingError(
                    f"スライド {number}（{title}）の原稿に合図〔{cue}〕がありますが、"
                    f"この画面の段は {total} 個です。\n"
                    "  段は画面に置かれた順に数えます(表は 1 行、図解は 1 項目、"
                    "文章は 1 段落、絵は `図@1.png` … の枚数)。"
                )
            image = render_step(base, elements, step)
            path = os.path.join(outdir, f"slide_{number:03d}.step{step:02d}.png")
            motion = _motion_of(elements, step)
            if motion:
                staged.motions[cue] = _render_motion(
                    base, elements, step, motion, os.path.splitext(path)[0]
                )
        image.save(path)
        staged.images[cue] = os.path.abspath(path)
    return staged


def render_step(base: Image.Image, elements: Sequence[Element], step: Optional[int]) -> Image.Image:
    """n 段目の画面を作る。`step` が None なら、すべて出した状態。"""
    return _render(base, elements, step, None)


def _render(
    base: Image.Image,
    elements: Sequence[Element],
    step: Optional[int],
    figure_frame: Optional[str],
) -> Image.Image:
    canvas = base.copy()
    pixels = np.asarray(canvas).astype(np.float32)
    pastes: List[Tuple[str, Box]] = []
    first = 1
    for element in elements:
        last = first + element.count - 1
        if step is not None and step < first:
            if element.frame:
                # 外枠は線が太く、角が丸い。縁のにじみまで覆うよう、広めに取る。
                _veil(pixels, element.frame, HIDDEN, pad=_PAD * 3)
            for part in element.parts or [element.box]:
                _veil(pixels, part, HIDDEN, ground=_ground(pixels, element.box))
        elif element.frames:
            if step is None or step > last:
                pastes.append((element.final, element.box))
            else:
                pastes.append((figure_frame or element.frames[step - first], element.box))
        elif step is not None and step > last:
            for part in element.parts or [element.box]:
                _veil(pixels, part, PAST, pad=0, ground=_ground(pixels, element.box))
        first = last + 1
    canvas = Image.fromarray(np.clip(pixels, 0, 255).astype(np.uint8), "RGB")
    for path, box in pastes:
        if not path:
            continue
        with Image.open(path) as picture:
            size = (box[2] - box[0], box[3] - box[1])
            canvas.paste(picture.convert("RGB").resize(size, Image.LANCZOS), box[:2])
    return canvas


def _motion_of(elements: Sequence[Element], step: int) -> List[str]:
    first = 1
    for element in elements:
        last = first + element.count - 1
        if first <= step <= last:
            return element.motions.get(step - first + 1, [])
        first = last + 1
    return []


def _render_motion(
    base: Image.Image, elements: Sequence[Element], step: int, frames: Sequence[str], stem: str
) -> List[str]:
    """絵の動き(連番)を、スライドの中に置いた形で書き出す。"""
    outdir = stem + ".motion"
    os.makedirs(outdir, exist_ok=True)
    paths: List[str] = []
    for position, frame in enumerate(frames):
        path = os.path.join(outdir, f"{position:03d}.png")
        _render(base, elements, step, frame).save(path, compress_level=1)
        paths.append(os.path.abspath(path))
    return paths


def _veil(
    pixels: np.ndarray,
    box: Box,
    amount: float,
    pad: int = _PAD,
    ground: Optional[np.ndarray] = None,
) -> None:
    """指定した場所を、まわりの地の色へ寄せる(その場で書き換える)。"""
    height, width = pixels.shape[:2]
    left, top = max(0, box[0] - pad), max(0, box[1] - pad)
    right, bottom = min(width, box[2] + pad), min(height, box[3] + pad)
    if right <= left or bottom <= top:
        return
    if ground is None:
        ground = _ground(pixels, (left, top, right, bottom))
    region = pixels[top:bottom, left:right]
    region += (ground - region) * amount


def _ground(pixels: np.ndarray, box: Box) -> np.ndarray:
    """その場所の左右すぐ外の色(地の色)。

    白とは限らない。枠図の項目は、枠の地の色の上に置かれている。上下ではなく
    左右を見るのは、表の行や縦に積んだ項目では、上下が隣の行だから。
    縁の線やにじみを拾わないよう、少し離れたところを見る。
    """
    left, top, right, bottom = box
    height, width = pixels.shape[:2]
    top, bottom = max(0, top), min(height, bottom)
    strips = []
    if left >= 9:
        strips.append(pixels[top:bottom, left - 9 : left - 5])
    if right + 9 <= width:
        strips.append(pixels[top:bottom, right + 5 : right + 9])
    if not strips:
        return np.array([255.0, 255.0, 255.0], dtype=np.float32)
    samples = np.concatenate([strip.reshape(-1, 3) for strip in strips])
    return np.median(samples, axis=0).astype(np.float32)


# ---------------------------------------------------------------------------
# 画面の上のものを数える
# ---------------------------------------------------------------------------

#: 項目ごとに分けて段にする図解。ほかの図解(境界・レーン・配置)は、項目どうしの
#: 関係が図の主題なので、ばらして出すと別の図に見える。まとめて 1 段にする。
_SPLIT_DIAGRAMS = (DIAGRAM_FLOW, DIAGRAM_FLOW_ACROSS, DIAGRAM_FRAME, DIAGRAM_STEPS)


def slide_elements(
    slide, base: Image.Image, scale: float, search_dirs: Sequence[str] = ()
) -> Tuple[List[Element], List[str]]:
    """1 枚の画面に置かれたものを、段にする単位で並べる(置かれた順)。"""
    warnings: List[str] = []
    pixels = np.asarray(base)
    title = slide.shapes.title
    elements: List[Element] = []
    diagram: Optional[Tuple[str, List[object], Box]] = None
    # 文章の箱は、文字の量より大きく取ってある(下の資料名・ページ番号に重なる
    # こともある)。段落を画像から読むときに、それらを文章の行として数えない。
    floor = min(
        [
            _box(shape, scale)[1]
            for shape in slide.shapes
            if parse_shape_name(getattr(shape, "name", ""))[0] == SHAPE_FOOTER
        ],
        default=base.height,
    )

    def clip(box: Box) -> Box:
        return box[0], box[1], box[2], min(box[3], floor - 2)

    def close_diagram() -> None:
        nonlocal diagram
        if diagram is None:
            return
        language, shapes, outer = diagram
        diagram = None
        if not shapes:
            return
        groups = _diagram_groups(shapes) if language in _SPLIT_DIAGRAMS else [shapes]
        for position, group in enumerate(groups):
            boxes = [_box(s, scale) for s in group]
            element = Element(_union(boxes), parts=boxes)
            if position == 0 and language == DIAGRAM_FRAME:
                # 枠図の外枠は、最初の項目と一緒に出す。空の枠だけが先に映ると、
                # 何かを読み込んでいる途中に見える。
                element.frame = outer
            elements.append(element)

    for shape in slide.shapes:
        if title is not None and shape.element is title.element:
            continue
        kind, _, language = parse_shape_name(getattr(shape, "name", ""))
        if kind in (SHAPE_FOOTER, SHAPE_LABEL):
            continue
        if kind == SHAPE_DIAGRAM:
            close_diagram()
            diagram = (language, [], _box(shape, scale))
            continue
        if kind == SHAPE_DIAGRAM_ITEM and diagram is not None:
            diagram[1].append(shape)
            continue
        close_diagram()
        if getattr(shape, "has_table", False):
            rows = len(shape.table.rows) - (1 if shape.table.first_row else 0)
            boxes = _table_rows(pixels, _box(shape, scale), len(shape.table.rows))
            if boxes is None:
                warnings.append("表の行を画像から読み取れなかったため、表をまとめて 1 段にしました。")
                elements.append(Element(_box(shape, scale)))
            else:
                elements.extend(Element(box) for box in boxes[len(boxes) - rows :])
        elif _is_picture(shape):
            elements.append(_figure_element(shape, scale, search_dirs))
        elif getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
            paragraphs = [p for p in shape.text_frame.paragraphs if p.text.strip()]
            area = clip(_box(shape, scale))
            boxes = _paragraph_boxes(pixels, area, len(paragraphs))
            if boxes is None:
                if len(paragraphs) > 1:
                    warnings.append(
                        "文章の段落を画像から読み取れなかったため、文章をまとめて 1 段にしました。"
                    )
                boxes = [_ink_box(pixels, area) or area]
            elements.extend(Element(box) for box in boxes)
    close_diagram()
    return elements, warnings


def _box(shape, scale: float) -> Box:
    left, top = int(round(shape.left * scale)), int(round(shape.top * scale))
    return left, top, left + int(round(shape.width * scale)), top + int(round(shape.height * scale))


def _union(boxes: Sequence[Box]) -> Box:
    return (
        min(b[0] for b in boxes),
        min(b[1] for b in boxes),
        max(b[2] for b in boxes),
        max(b[3] for b in boxes),
    )


def _is_picture(shape) -> bool:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    try:
        return shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    except (AttributeError, ValueError, KeyError):
        return False


def _has_text(shape) -> bool:
    return bool(getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip())


def _diagram_groups(shapes: Sequence[object]) -> List[List[object]]:
    """図解の図形を、項目ごとにまとめる。

    図形は「項目 -> つなぐもの(矢印・蹴込み) -> 項目 …」の順に置かれている。
    文字の無い図形はつなぐもので、**次の項目に付ける**(矢印は、行き先と
    一緒に出るのが自然)。同じ行に並んだ文字の図形(階段図の札と中身)は、
    1 つの項目として扱う。
    """
    groups: List[List[object]] = []
    current: List[object] = []
    for shape in shapes:
        texts = [s for s in current if _has_text(s)]
        if not _has_text(shape):
            if texts:
                groups.append(current)
                current = []
            current.append(shape)
        elif not texts or _same_row(texts[-1], shape):
            current.append(shape)
        else:
            groups.append(current)
            current = [shape]
    if current:
        groups.append(current)
    return groups


def _same_row(a, b) -> bool:
    """同じ行に、横に続けて置かれているか(階段図の札と中身)。"""
    overlap = min(a.top + a.height, b.top + b.height) - max(a.top, b.top)
    touching = abs((a.left + a.width) - b.left) <= max(a.height, b.height) * 0.1
    return overlap > min(a.height, b.height) * 0.9 and touching


def _table_rows(pixels: np.ndarray, box: Box, count: int) -> Optional[List[Box]]:
    """表の行の場所を、画像の色の切れ目から読み取る。

    行の高さは資料に書かれていない(文字が折り返すと、描く側が行を伸ばす)。
    表の左端のすぐ内側には文字が無いので、そこを縦になぞって、色が変わる
    ところを行の境目とする。行の数が合わなければ、読み取れなかったとして返す。
    """
    left, top, right, _ = box
    height, width = pixels.shape[:2]
    x0, x1 = min(width - 1, left + 3), min(width, left + 9)
    if x1 <= x0:
        return None
    column = pixels[:, x0:x1].astype(np.float32).mean(axis=1)
    ground = _ground(pixels.astype(np.float32), (left, max(0, top), right, min(height, top + 40)))

    runs: List[List[int]] = []  # [始まり, 終わり)
    y = max(0, top - 2)
    start, color = None, None
    while y < height:
        here = column[y]
        if start is None:
            if np.abs(here - ground).max() > 5:
                start, color = y, here
        elif np.abs(here - color).max() > 6:
            runs.append([start, y])
            start, color = (y, here) if np.abs(here - ground).max() > 5 else (None, None)
            if start is None and len([r for r in runs if r[1] - r[0] >= 6]) >= count:
                break
        y += 1
    if start is not None:
        runs.append([start, height])
    rows = [run for run in runs if run[1] - run[0] >= 6]  # 細い線は罫線
    if len(rows) != count:
        return None
    return [(left, a, right, b) for a, b in rows]


def _ink_rows(pixels: np.ndarray, box: Box) -> List[Tuple[int, int]]:
    """その場所の中で、文字のある横帯(行)を上から並べる。"""
    left, top, right, bottom = box
    height, width = pixels.shape[:2]
    left, right = max(0, left), min(width, right)
    top, bottom = max(0, top), min(height, bottom)
    if right <= left or bottom <= top:
        return []
    region = pixels[top:bottom, left:right].astype(np.float32)
    ground = np.median(region.reshape(-1, 3), axis=0)
    ink = (np.abs(region - ground).max(axis=2) > 48).any(axis=1)
    bands: List[Tuple[int, int]] = []
    start = None
    for y, on in enumerate(ink):
        if on and start is None:
            start = y
        elif not on and start is not None:
            bands.append((top + start, top + y))
            start = None
    if start is not None:
        bands.append((top + start, bottom))
    # 濁点や句読点だけの細い帯を、近くの行にまとめる。
    merged: List[Tuple[int, int]] = []
    for band in bands:
        if merged and band[0] - merged[-1][1] <= 3:
            merged[-1] = (merged[-1][0], band[1])
        else:
            merged.append(band)
    return merged


def _ink_box(pixels: np.ndarray, box: Box) -> Optional[Box]:
    rows = _ink_rows(pixels, box)
    if not rows:
        return None
    return box[0], rows[0][0], box[2], rows[-1][1]


def _paragraph_boxes(pixels: np.ndarray, box: Box, count: int) -> Optional[List[Box]]:
    """文章の段落の場所を、行のすき間から読み取る。

    段落のあいだは、行のあいだより広く空けて描かれる。行のすき間のうち広い
    ものから `段落の数 - 1` 個を段落の切れ目とする。切れ目とそうでないすき間の
    差がはっきりしないときは、読み取れなかったとして返す。
    """
    if count <= 0:
        return None
    rows = _ink_rows(pixels, box)
    if len(rows) < count:
        return None
    if count == 1:
        return [(box[0], rows[0][0], box[2], rows[-1][1])]
    gaps = [(rows[i + 1][0] - rows[i][1], i) for i in range(len(rows) - 1)]
    ranked = sorted(gaps, reverse=True)
    if len(ranked) > count - 1 and ranked[count - 2][0] - ranked[count - 1][0] < 3:
        return None
    breaks = sorted(index for _, index in ranked[: count - 1])
    boxes: List[Box] = []
    first = 0
    for index in breaks + [len(rows) - 1]:
        boxes.append((box[0], rows[first][0], box[2], rows[index][1]))
        first = index + 1
    return boxes


# ---------------------------------------------------------------------------
# 絵の段
# ---------------------------------------------------------------------------

_FRAME_NAME = re.compile(r"^\d+\.png$", re.IGNORECASE)


def step_path(image_path: str, step: int) -> str:
    """絵 `図.png` の n 段目を置く場所(`図@n.png`)。"""
    stem, suffix = os.path.splitext(image_path)
    return f"{stem}@{step}{suffix}"


def motion_dir(image_path: str, step: int) -> str:
    """絵 `図.png` の n 段目に入るときの動き(連番)を置く場所(`図@n/`)。"""
    return f"{os.path.splitext(image_path)[0]}@{step}"


def figure_steps(image_path: str) -> List[str]:
    """絵の隣にある段の画像を、1 段目から途切れるまで集める。"""
    found: List[str] = []
    while os.path.isfile(step_path(image_path, len(found) + 1)):
        found.append(step_path(image_path, len(found) + 1))
    return found


def _figure_element(shape, scale: float, search_dirs: Sequence[str]) -> Element:
    element = Element(_box(shape, scale), kind="figure")
    source = _figure_source(shape, search_dirs)
    if not source:
        return element
    element.final = source
    element.frames = figure_steps(source)
    for step in range(1, len(element.frames) + 1):
        directory = motion_dir(source, step)
        if os.path.isdir(directory):
            names = sorted(name for name in os.listdir(directory) if _FRAME_NAME.match(name))
            if names:
                element.motions[step] = [os.path.join(directory, name) for name in names]
    if not element.frames:
        element.kind = "item"
    return element


def _figure_source(shape, search_dirs: Sequence[str]) -> str:
    """絵のもとのファイル。資料を作るときに、図形の説明へ場所を書いてある。"""
    try:
        described = shape._element.nvPicPr.cNvPr.get("descr") or ""
    except AttributeError:
        return ""
    if not described:
        return ""
    candidates = [described] if os.path.isabs(described) else []
    for directory in search_dirs:
        candidates.append(os.path.join(directory, described))
        candidates.append(os.path.join(directory, os.path.basename(described)))
    for candidate in candidates:
        if os.path.isfile(candidate):
            return os.path.abspath(candidate)
    return ""


# ---------------------------------------------------------------------------
# 一覧
# ---------------------------------------------------------------------------


def _title_of(slide) -> str:
    title = slide.shapes.title
    return title.text.strip() if title is not None else ""


def _slide_image_paths(slides_dir: str) -> Dict[int, str]:
    path = os.path.join(slides_dir, images_mod.MANIFEST_NAME)
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    return {
        int(entry["index"]): os.path.join(slides_dir, str(entry["file"]))
        for entry in data.get("slides", [])
    }


def _write_manifest(result: StagingResult) -> str:
    data = {
        "count": len(result.slides),
        "slides": [slide.to_dict() for slide in result.slides],
        "warnings": result.warnings,
    }
    path = os.path.join(result.outdir, MANIFEST_NAME)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write("\n")
    return os.path.abspath(path)


def load_steps(slides_dir: Optional[str]) -> Dict[int, StagedSlide]:
    """書き出してある段の一覧を読む。無ければ空(これまでどおりの動画になる)。"""
    if not slides_dir:
        return {}
    outdir = os.path.join(slides_dir, STEPS_DIR)
    path = os.path.join(outdir, MANIFEST_NAME)
    if not os.path.isfile(path):
        return {}
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    staged: Dict[int, StagedSlide] = {}
    for entry in data.get("slides", []):
        slide = StagedSlide(index=int(entry["index"]), count=int(entry.get("count", 0)))
        slide.images = {
            key: os.path.join(outdir, name) for key, name in entry.get("images", {}).items()
        }
        slide.motions = {
            key: [os.path.join(outdir, name) for name in names]
            for key, names in entry.get("motions", {}).items()
        }
        staged[slide.index] = slide
    return staged


# ---------------------------------------------------------------------------
# 時間割り
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class Shot:
    """1 枚の中の 1 場面。`at` はそのスライドが出てからの秒。"""

    at: float
    image: str
    cue: str = ""
    motion: Tuple[str, ...] = ()

    def to_dict(self) -> dict:
        data = {"at": round(self.at, 3), "cue": self.cue, "image": os.path.basename(self.image)}
        if self.motion:
            data["motion_frames"] = len(self.motion)
        return data


def slide_shots(
    base_image: str,
    staged: Optional[StagedSlide],
    cues: Sequence[dict],
    lead: float = 0.0,
) -> List[Shot]:
    """1 枚の中で、何秒目にどの画像へ変えるかを並べる。

    画面は、合図の文が読まれる少し手前(`lead` 秒)で変える。声より先に画面が
    変わっていたほうが、「いま言っているもの」を探さずに済む。最初の文に合図が
    あれば、その画面から始める(一瞬だけ全体が映るのを避ける)。
    """
    shots = [Shot(0.0, base_image, CUE_ALL)]
    if staged is None:
        return shots
    if staged.images.get(CUE_START):
        shots = [Shot(0.0, staged.images[CUE_START], CUE_START)]
    for cue in cues:
        name = str(cue.get("step", ""))
        image = staged.images.get(name, base_image if name == CUE_ALL else "")
        if not image:
            continue
        at = max(0.0, float(cue.get("at", 0.0)) - lead)
        shot = Shot(at, image, name, tuple(staged.motions.get(name, ())))
        if cue.get("first") or at <= 0.0:
            shots = [Shot(0.0, image, name)]  # 始まりの画面に、動きは付けない
        elif at > shots[-1].at:
            shots.append(shot)
        else:
            shots[-1] = Shot(shots[-1].at, image, name, shot.motion)
    return shots
