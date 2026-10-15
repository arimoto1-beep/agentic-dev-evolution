"""画面いっぱいの場面(`レイアウト: 全画面`)。

動画の 1 場面として描いた絵を、スライドの枠(見出し・罫線・資料名・ページ番号)に
入れずに、そのまま画面全体に映すためのレイアウト。崩れると困るのは次のところ:

* 絵が画面全体を覆い、資料名とページ番号が置かれないこと
* 見出しが資料の中には残ること(章立てと、資料を開いた人のため)
* 絵 1 枚以外を書いたら、書いた時点で止まること
* 段の絵と動きの連番が、置き直されずにそのまま動画のコマになること
* 最初の合図までの姿(`図@0.png`)が使われること(覆うと白紙になる)
"""

import json
import os

import pytest
from PIL import Image
from pptx import Presentation

from note2slides import staging
from note2slides.model import KIND_FULL
from note2slides.renderer import render_deck
from note2slides.scenario import ScenarioError, build_deck, parse_scenario

HEAD = "---\ntype: scenario\ntitle: 全画面のテスト\n---\n\n"
BODY = "## 場面の題\n\n### 設定\n- レイアウト: 全画面\n\n### 画面\n\n![絵の説明](figures/scene.png)\n\n### ナレーション\n\n〔1〕一つ目です。〔2〕二つ目です。\n"
SIZE = (1920, 1080)


def make_scene(tmp_path, steps=2, opening=True, motion_for=None, size=SIZE):
    """場面の絵と、その段(色違い)を作る。"""
    folder = tmp_path / "figures"
    folder.mkdir(exist_ok=True)
    path = str(folder / "scene.png")
    Image.new("RGB", size, (10, 10, 10)).save(path)
    if opening:
        Image.new("RGB", size, (0, 0, 90)).save(staging.step_path(path, 0))
    for step in range(1, steps + 1):
        Image.new("RGB", size, (step * 60, 0, 0)).save(staging.step_path(path, step))
    if motion_for:
        os.makedirs(staging.motion_dir(path, motion_for))
        for index in range(3):
            Image.new("RGB", size, (0, 80 * (index + 1), 0)).save(
                os.path.join(staging.motion_dir(path, motion_for), f"{index:03d}.png")
            )
    return path


def make_deck(tmp_path, body=BODY):
    source = tmp_path / "lesson.md"
    deck = build_deck(parse_scenario(HEAD + body, source_path=str(source)))
    path = str(tmp_path / "lesson.pptx")
    render_deck(deck, path)
    slides_dir = tmp_path / "slides"
    slides_dir.mkdir(exist_ok=True)
    Image.new("RGB", SIZE, (255, 255, 255)).save(str(slides_dir / "slide_001.png"))
    with open(slides_dir / "slides.json", "w", encoding="utf-8") as f:
        json.dump({"slides": [{"index": 1, "file": "slide_001.png"}]}, f)
    return deck, path, str(slides_dir)


# ---------------------------------------------------------------------------
# シナリオと資料
# ---------------------------------------------------------------------------


def test_full_layout_makes_a_slide_that_is_only_the_picture(tmp_path):
    make_scene(tmp_path)
    deck, path, _ = make_deck(tmp_path)
    assert deck.slides[0].kind == KIND_FULL

    presentation = Presentation(path)
    slide = presentation.slides[0]
    pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
    assert len(pictures) == 1
    picture = pictures[0]
    assert (picture.left, picture.top) == (0, 0)
    assert picture.width == presentation.slide_width
    assert abs(picture.height - presentation.slide_height) <= 1
    # 資料名・ページ番号・札は置かない(絵が画面全体を持つ)。
    names = [shape.name for shape in slide.shapes]
    assert not any("footer" in name or "label" in name for name in names)


def test_full_layout_keeps_the_heading_in_the_deck_under_the_picture(tmp_path):
    make_scene(tmp_path)
    _, path, _ = make_deck(tmp_path)
    slide = Presentation(path).slides[0]
    assert slide.shapes.title.text == "場面の題"
    order = [shape.shape_type == 13 for shape in slide.shapes]
    assert order.index(True) > list(slide.shapes).index(slide.shapes.title)  # 絵が上


def test_full_layout_keeps_the_narration_in_the_notes(tmp_path):
    make_scene(tmp_path)
    _, path, _ = make_deck(tmp_path)
    assert "一つ目です。" in Presentation(path).slides[0].notes_slide.notes_text_frame.text


def test_picture_taller_than_the_screen_is_fitted_not_stretched(tmp_path):
    make_scene(tmp_path, steps=0, opening=False, size=(1000, 1000))
    _, path, _ = make_deck(tmp_path)
    presentation = Presentation(path)
    picture = [s for s in presentation.slides[0].shapes if s.shape_type == 13][0]
    assert abs(picture.height - presentation.slide_height) <= 1
    assert abs(picture.width - picture.height) <= 1
    assert abs(picture.left - (presentation.slide_width - picture.width) / 2) <= 1


@pytest.mark.parametrize(
    "screen",
    ["文章だけです。", "![](figures/scene.png)\n\n補足の文です。", "![](figures/scene.png)\n\n![](figures/scene.png)", ""],
)
def test_full_layout_refuses_anything_but_one_picture(tmp_path, screen):
    make_scene(tmp_path)
    body = f"## 場面\n\n### 設定\n- レイアウト: 全画面\n\n### 画面\n\n{screen}\n"
    with pytest.raises(ScenarioError) as error:
        build_deck(parse_scenario(HEAD + body, source_path=str(tmp_path / "lesson.md")))
    assert "全画面" in str(error.value) and "絵 1 枚" in str(error.value)


def test_full_layout_refuses_a_label_and_says_where_to_draw_it(tmp_path):
    make_scene(tmp_path)
    body = "## 場面\n\n### 設定\n- レイアウト: 全画面\n- 札: 推測\n\n### 画面\n\n![](figures/scene.png)\n"
    with pytest.raises(ScenarioError) as error:
        parse_scenario(HEAD + body, source_path=str(tmp_path / "lesson.md"))
    assert "札" in str(error.value) and "絵の中" in str(error.value)


def test_english_name_can_be_used(tmp_path):
    make_scene(tmp_path)
    body = "## Scene\n\n### settings\n- layout: full\n\n### screen\n\n![](figures/scene.png)\n"
    deck = build_deck(parse_scenario(HEAD + body, source_path=str(tmp_path / "lesson.md")))
    assert deck.slides[0].kind == KIND_FULL


# ---------------------------------------------------------------------------
# 段
# ---------------------------------------------------------------------------


def color_of(path):
    with Image.open(path) as image:
        return image.convert("RGB").getpixel((960, 540))


def test_step_pictures_and_motion_frames_are_used_as_they_are(tmp_path):
    scene = make_scene(tmp_path, steps=2, motion_for=2)
    _, path, slides_dir = make_deck(tmp_path)
    result = staging.build_steps(path, slides_dir, {1: ["1", "2", "全体"]})

    (slide,) = result.slides
    assert slide.count == 2
    # 絵そのものを指す(スライド画像の上へ置き直した写しを作らない)。
    assert slide.images["1"] == staging.step_path(scene, 1)
    assert slide.images["2"] == staging.step_path(scene, 2)
    assert slide.images["全体"] == scene
    assert [os.path.dirname(p) for p in slide.motions["2"]] == [staging.motion_dir(scene, 2)] * 3
    assert not [name for name in os.listdir(result.outdir) if name.endswith(".png")]


def test_opening_picture_is_the_screen_before_the_first_cue(tmp_path):
    scene = make_scene(tmp_path, steps=2, opening=True)
    _, path, slides_dir = make_deck(tmp_path)
    result = staging.build_steps(path, slides_dir, {1: ["1", "2"]})
    assert result.slides[0].images["0"] == staging.step_path(scene, 0)
    assert color_of(result.slides[0].images["0"]) == (0, 0, 90)


def test_without_an_opening_picture_the_screen_before_the_first_cue_is_covered(tmp_path):
    make_scene(tmp_path, steps=2, opening=False)
    _, path, slides_dir = make_deck(tmp_path)
    result = staging.build_steps(path, slides_dir, {1: ["1", "2"]})
    assert color_of(result.slides[0].images["0"]) == (255, 255, 255)


def test_manifest_points_back_to_the_pictures(tmp_path):
    scene = make_scene(tmp_path, steps=2, motion_for=2)
    _, path, slides_dir = make_deck(tmp_path)
    staging.build_steps(path, slides_dir, {1: ["1", "2", "全体"]})

    loaded = staging.load_steps(slides_dir)[1]
    assert loaded.images["1"] == staging.step_path(scene, 1)
    assert all(os.path.isfile(p) for p in loaded.images.values())
    assert all(os.path.isfile(p) for p in loaded.motions["2"])


def test_step_pictures_of_another_size_are_placed_on_the_slide_as_before(tmp_path):
    make_scene(tmp_path, steps=2, opening=False, size=(960, 540))
    _, path, slides_dir = make_deck(tmp_path)
    result = staging.build_steps(path, slides_dir, {1: ["1", "2"]})
    step = result.slides[0].images["1"]
    assert os.path.dirname(step) == result.outdir  # 動画の大きさに置き直した写し
    with Image.open(step) as image:
        assert image.size == SIZE
    assert color_of(step) == (60, 0, 0)


def test_opening_picture_is_also_used_for_a_figure_inside_a_slide(tmp_path):
    # 全画面でなくても、絵の「まだ何も出していない姿」を絵の側で用意できる。
    folder = tmp_path / "figures"
    folder.mkdir()
    figure = str(folder / "scene.png")
    Image.new("RGB", (240, 105), (10, 10, 10)).save(figure)
    Image.new("RGB", (240, 105), (0, 0, 90)).save(staging.step_path(figure, 0))
    Image.new("RGB", (240, 105), (60, 0, 0)).save(staging.step_path(figure, 1))
    body = "## 絵\n\n### 画面\n\n![](figures/scene.png)\n"
    _, path, slides_dir = make_deck(tmp_path, body)
    result = staging.build_steps(path, slides_dir, {1: ["1"]})
    with Image.open(result.slides[0].images["0"]) as image:
        assert image.convert("RGB").getpixel((960, 600)) == (0, 0, 90)
