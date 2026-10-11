"""1 枚の画面を段に分ける(`staging`)。

段は「ナレーションがそこまで進んだら出す」ものなので、**数え方が画面と
合っていること** がいちばん大事になる。数え違えると、話していない行が先に出たり、
話している行が隠れたままになったりして、しかも通しで見るまで気付けない。

* 図解・表・文章・絵が、置かれた順に、決めた単位で数えられること
* n 段目の画面で、先の段が見えず、いまの段がそのまま、前の段が薄くなること
* 絵の段(`図@1.png` …)と動きが、絵の場所に置かれること
* 合図が指す段が無いときに、動画にせずに止まること
* 何秒目にどの画面へ変えるかが、合図の時刻から決まること

LibreOffice は呼ばない。資料は実際に作り、スライド画像は同じ大きさの絵で代える
(図形の場所は資料から読むので、画像の中身が無くても数え方は確かめられる)。
"""

import json
import os

import numpy as np
import pytest
from PIL import Image

from note2slides import staging
from note2slides.renderer import render_deck
from note2slides.scenario import build_deck, parse_scenario
from note2slides.staging import Element, Shot, StagedSlide, StagingError

HEAD = "---\ntype: scenario\ntitle: 段のテスト\n---\n\n"
WHITE = (255, 255, 255)
SIZE = (1920, 1080)


def make_deck(tmp_path, body):
    """シナリオから資料を作り、同じ大きさの白いスライド画像を添える。"""
    source = tmp_path / "lesson.md"
    deck = build_deck(parse_scenario(HEAD + body, source_path=str(source)))
    path = str(tmp_path / "lesson.pptx")
    render_deck(deck, path)
    slides_dir = tmp_path / "slides"
    slides_dir.mkdir(exist_ok=True)
    entries = []
    for index in range(1, len(deck.slides) + 1):
        name = f"slide_{index:03d}.png"
        Image.new("RGB", SIZE, WHITE).save(str(slides_dir / name))
        entries.append({"index": index, "file": name})
    with open(slides_dir / "slides.json", "w", encoding="utf-8") as f:
        json.dump({"slides": entries}, f)
    return path, str(slides_dir)


def elements_of(tmp_path, body, image=None, slide_number=1):
    from pptx import Presentation

    path, _ = make_deck(tmp_path, body)
    presentation = Presentation(path)
    slide = list(presentation.slides)[slide_number - 1]
    base = image or Image.new("RGB", SIZE, WHITE)
    return staging.slide_elements(
        slide, base, base.width / presentation.slide_width, [str(tmp_path)]
    )


# ---------------------------------------------------------------------------
# 数え方
# ---------------------------------------------------------------------------


def test_flow_items_are_one_step_each_and_the_arrow_comes_with_its_target(tmp_path):
    elements, _ = elements_of(
        tmp_path, "## 流れ\n\n### 画面\n\n```流れ\n一つ目\n二つ目\n三つ目\n```\n"
    )
    assert len(elements) == 3
    assert [len(e.parts) for e in elements] == [1, 2, 2]  # 2 つ目からは、矢印 + 箱
    tops = [e.box[1] for e in elements]
    assert tops == sorted(tops)


def test_frame_items_are_one_step_each_and_the_frame_itself_is_not_a_step(tmp_path):
    elements, _ = elements_of(tmp_path, "## 枠\n\n### 画面\n\n```枠\n一つ目\n二つ目\n三つ目\n```\n")
    assert len(elements) == 3
    assert all(len(e.parts) == 1 for e in elements)


def test_frame_itself_appears_with_its_first_item(tmp_path):
    """空の枠だけが先に映ると、読み込み中に見える。何も出していない段では枠も隠す。"""
    elements, _ = elements_of(tmp_path, "## 枠\n\n### 画面\n\n```枠\n一つ目\n二つ目\n```\n")
    first, second = elements
    assert first.frame is not None and second.frame is None
    left, top, right, bottom = first.frame
    assert left < first.box[0] and top < first.box[1] and bottom > second.box[3]

    base = Image.new("RGB", SIZE, WHITE)
    pixels = np.asarray(base).copy()
    pixels[top:bottom, left:right] = (244, 245, 247)  # 枠の地
    for element in elements:
        l, t, r, b = element.box
        pixels[t:b, l:r] = (20, 40, 60)
    base = Image.fromarray(pixels)
    inside = ((left + first.box[0]) // 2, (top + bottom) // 2)  # 枠の中で、項目の外
    assert staging.render_step(base, elements, 0).getpixel(inside) == WHITE
    assert staging.render_step(base, elements, 1).getpixel(inside) == (244, 245, 247)
    assert staging.render_step(base, elements, 2).getpixel(inside) == (244, 245, 247)


def test_step_diagram_counts_levels_and_the_reach_mark(tmp_path):
    body = (
        "## 階段\n\n### 画面\n\n```階段\n"
        "一段目: はじめ\n二段目: つぎ\n← ここまで届いた\n三段目: その先\n```\n"
    )
    elements, _ = elements_of(tmp_path, body)
    # 段 1(札 + 中身) / 段 2(蹴込み + 札 + 中身) / 到達点(矢印 + 札) / 段 3
    assert [len(e.parts) for e in elements] == [2, 3, 2, 3]


def test_other_diagrams_are_one_step_as_a_whole(tmp_path):
    body = "## 境界\n\n### 画面\n\n```境界\n上の側\n↓ 降りるもの\n---\n下の側\n```\n"
    elements, _ = elements_of(tmp_path, body)
    assert len(elements) == 1


def test_text_after_a_diagram_is_its_own_step(tmp_path):
    body = "## 流れと文\n\n### 画面\n\n```流れ\n一つ目\n二つ目\n```\n\nまとめの文です。\n"
    elements, _ = elements_of(tmp_path, body)
    assert len(elements) == 3


def test_title_label_and_footer_are_not_steps(tmp_path):
    body = "## 見出し\n\n### 設定\n\n- 札: 推測\n\n### 画面\n\n```枠\nただ一つ\n```\n"
    elements, _ = elements_of(tmp_path, body)
    assert len(elements) == 1


def banded(rows, left=100, right=900, top=200, heights=(40, 60, 90, 50)):
    """行ごとに色が変わる表のような画像。行の境目(下端)も返す。"""
    pixels = np.full((600, 1000, 3), 255, dtype=np.uint8)
    fills = [(79, 129, 189), (208, 216, 232), (233, 237, 244)]
    y = top
    edges = []
    for row in range(rows):
        height = heights[row % len(heights)]
        pixels[y : y + height, left:right] = fills[0] if row == 0 else fills[1 + (row % 2)]
        pixels[y + height - 1, left:right] = 255  # 行のあいだの白い罫線
        y += height
        edges.append(y)
    return pixels, (left, top, right, top + 30), edges


def test_table_rows_are_read_from_the_picture_even_when_rows_have_grown():
    """行の高さは資料に書かれていない。文字が折り返して伸びた行も、色の境目から読む。"""
    pixels, box, edges = banded(4)
    rows = staging._table_rows(pixels, box, 4)
    assert rows is not None and len(rows) == 4
    assert [abs(row[3] - edge) <= 1 for row, edge in zip(rows, edges)] == [True] * 4


def test_table_rows_give_up_when_the_count_does_not_match():
    pixels, box, _ = banded(3)
    assert staging._table_rows(pixels, box, 5) is None


def lines(tops, height=20, gap_color=255):
    pixels = np.full((400, 600, 3), gap_color, dtype=np.uint8)
    for top in tops:
        pixels[top : top + height, 50:500] = 30
    return pixels


def test_paragraphs_are_split_at_the_widest_gaps_between_lines():
    # 3 段落: 1 行 / 2 行 / 1 行。段落のあいだ(30)は、行のあいだ(10)より広い。
    pixels = lines([50, 100, 130, 180])
    boxes = staging._paragraph_boxes(pixels, (40, 40, 560, 380), 3)
    assert [(box[1], box[3]) for box in boxes] == [(50, 70), (100, 150), (180, 200)]


def test_paragraphs_give_up_when_the_gaps_do_not_tell_them_apart():
    pixels = lines([50, 80, 110, 140])  # どのすき間も同じ
    assert staging._paragraph_boxes(pixels, (40, 40, 560, 380), 2) is None


# ---------------------------------------------------------------------------
# n 段目の画面
# ---------------------------------------------------------------------------


def three_bars():
    """地の色の上に、濃い帯が 3 本ある画像と、その 3 本を段にしたもの。"""
    base = Image.new("RGB", (300, 300), (240, 244, 250))
    pixels = np.asarray(base).copy()
    boxes = [(50, 40, 250, 80), (50, 120, 250, 160), (50, 200, 250, 240)]
    for left, top, right, bottom in boxes:
        pixels[top:bottom, left:right] = (20, 40, 60)
    return Image.fromarray(pixels), [Element(box) for box in boxes]


def middle(image, box):
    return image.getpixel(((box[0] + box[2]) // 2, (box[1] + box[3]) // 2))


def test_step_shows_the_current_item_hides_later_ones_and_recedes_earlier_ones():
    base, elements = three_bars()
    shown = staging.render_step(base, elements, 2)
    ground, ink = (240, 244, 250), (20, 40, 60)
    assert middle(shown, elements[1].box) == ink  # いまの段は、そのまま
    assert middle(shown, elements[2].box) == ground  # 先の段は、地の色(見えない)
    earlier = middle(shown, elements[0].box)
    assert ink[0] < earlier[0] < ground[0]  # 前の段は、薄くなるが残る


def test_step_zero_shows_nothing_and_none_shows_everything():
    base, elements = three_bars()
    nothing = staging.render_step(base, elements, 0)
    assert all(middle(nothing, e.box) == (240, 244, 250) for e in elements)
    everything = staging.render_step(base, elements, None)
    assert everything.tobytes() == base.tobytes()


def test_hiding_a_group_does_not_erase_what_lies_between_its_parts():
    """まとまりを 1 つの四角で覆うと、あいだにある別の図形まで消える。"""
    base, elements = three_bars()
    outer = Element(
        (50, 40, 250, 240), parts=[elements[0].box, elements[2].box]
    )  # 1 本目と 3 本目のまとまり
    shown = staging.render_step(base, [elements[1], outer], 1)
    assert middle(shown, elements[1].box) == (20, 40, 60)  # あいだの 2 本目は残る
    assert middle(shown, elements[0].box) == (240, 244, 250)


# ---------------------------------------------------------------------------
# 絵の段
# ---------------------------------------------------------------------------

FIGURE_BODY = "## 絵\n\n### 画面\n\n![](figures/part.png)\n"


def make_figure(tmp_path, steps=2, motion_for=None):
    """絵と、その段(色違い)を作る。"""
    folder = tmp_path / "figures"
    folder.mkdir(exist_ok=True)
    path = str(folder / "part.png")
    Image.new("RGB", (240, 105), (10, 10, 10)).save(path)
    for step in range(1, steps + 1):
        Image.new("RGB", (240, 105), (step * 60, 0, 0)).save(staging.step_path(path, step))
    if motion_for:
        os.makedirs(staging.motion_dir(path, motion_for))
        for index in range(3):
            Image.new("RGB", (240, 105), (0, 80 * (index + 1), 0)).save(
                os.path.join(staging.motion_dir(path, motion_for), f"{index:03d}.png")
            )
    return path


def test_figure_counts_as_many_steps_as_it_has_step_pictures(tmp_path):
    make_figure(tmp_path, steps=3)
    elements, _ = elements_of(tmp_path, FIGURE_BODY)
    assert [e.count for e in elements] == [3]


def test_figure_without_step_pictures_is_one_step(tmp_path):
    make_figure(tmp_path, steps=0)
    elements, _ = elements_of(tmp_path, FIGURE_BODY)
    assert [e.count for e in elements] == [1]


def test_build_steps_puts_the_step_picture_where_the_figure_is(tmp_path):
    make_figure(tmp_path, steps=2, motion_for=2)
    path, slides_dir = make_deck(tmp_path, FIGURE_BODY)
    result = staging.build_steps(path, slides_dir, {1: ["1", "2", "全体"]})

    (slide,) = result.slides
    assert slide.count == 2
    assert set(slide.images) == {"0", "1", "2", "全体"}
    elements, _ = elements_of(tmp_path, FIGURE_BODY)
    box = elements[0].box
    for cue, color in [("0", WHITE), ("1", (60, 0, 0)), ("2", (120, 0, 0)), ("全体", (10, 10, 10))]:
        with Image.open(slide.images[cue]) as image:
            assert middle(image.convert("RGB"), box) == color, cue
    assert len(slide.motions["2"]) == 3
    with Image.open(slide.motions["2"][0]) as frame:
        assert middle(frame.convert("RGB"), box) == (0, 80, 0)

    # 書き出した一覧から、同じものを読み戻せる(動画はこちらを読む)。
    loaded = staging.load_steps(slides_dir)[1]
    assert set(loaded.images) == set(slide.images)
    assert all(os.path.isfile(p) for p in loaded.images.values())
    assert all(os.path.isfile(p) for p in loaded.motions["2"])


def test_cue_beyond_the_last_step_stops_with_the_slide_and_the_count(tmp_path):
    body = "## 三つの箱\n\n### 画面\n\n```枠\n一つ目\n二つ目\n三つ目\n```\n"
    path, slides_dir = make_deck(tmp_path, body)
    with pytest.raises(StagingError) as error:
        staging.build_steps(path, slides_dir, {1: ["1", "4"]})
    message = str(error.value)
    assert "スライド 1" in message and "三つの箱" in message
    assert "〔4〕" in message and "3 個" in message


def test_steps_from_the_previous_run_do_not_remain(tmp_path):
    body = "## 箱\n\n### 画面\n\n```枠\n一つ目\n二つ目\n```\n"
    path, slides_dir = make_deck(tmp_path, body)
    staging.build_steps(path, slides_dir, {1: ["1", "2"]})
    assert staging.load_steps(slides_dir)
    staging.build_steps(path, slides_dir, {})  # 合図を消して作り直した
    assert staging.load_steps(slides_dir) == {}


# ---------------------------------------------------------------------------
# 時間割り
# ---------------------------------------------------------------------------


def staged(*cues, motions=None):
    slide = StagedSlide(index=1, count=3)
    slide.images = {cue: f"step{cue}.png" for cue in cues}
    slide.motions = motions or {}
    return slide


def test_slide_without_steps_is_one_shot_of_the_slide_itself():
    assert staging.slide_shots("base.png", None, []) == [Shot(0.0, "base.png", "全体")]


def test_slide_with_steps_starts_empty_until_the_first_cue():
    """出揃った画面を先に見せてから隠すと、これから話すことを先に読まれる。"""
    shots = staging.slide_shots(
        "base.png", staged("0", "1", "2"), [{"step": "1", "at": 3.0}, {"step": "2", "at": 6.0}]
    )
    assert [(s.at, s.image) for s in shots] == [(0.0, "step0.png"), (3.0, "step1.png"), (6.0, "step2.png")]


def test_cue_on_the_first_sentence_is_the_opening_screen():
    shots = staging.slide_shots(
        "base.png",
        staged("0", "1", "2"),
        [{"step": "1", "at": 0.3, "first": True}, {"step": "2", "at": 4.0}],
    )
    assert [(s.at, s.image) for s in shots] == [(0.0, "step1.png"), (4.0, "step2.png")]


def test_screen_changes_a_little_before_the_sentence_is_spoken():
    shots = staging.slide_shots(
        "base.png", staged("0", "1"), [{"step": "1", "at": 3.0}], lead=0.2
    )
    assert shots[1].at == pytest.approx(2.8)


def test_cue_for_everything_returns_to_the_slide_itself():
    shots = staging.slide_shots(
        "base.png", staged("0", "1"), [{"step": "1", "at": 2.0}, {"step": "全体", "at": 5.0}]
    )
    assert shots[-1].image == "base.png"


def test_motion_is_attached_to_the_shot_of_its_step():
    shots = staging.slide_shots(
        "base.png",
        staged("0", "1", "2", motions={"2": ["m0.png", "m1.png"]}),
        [{"step": "1", "at": 2.0}, {"step": "2", "at": 5.0}],
    )
    assert shots[2].motion == ("m0.png", "m1.png")
    assert shots[1].motion == ()
